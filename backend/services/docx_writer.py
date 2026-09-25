"""
Renders a parsed research brief into a .docx.

The parser (services/document_common.py), the section numbering, and the
design tokens (services/document_theme.py) are all shared with the PDF
writer, so the two formats of one brief agree on structure, numbering, and
references. Only the drawing differs.

Word-specific choices, each deliberate:

* Fonts are referenced by name, never embedded, so the body face has to be
  one the reader's Word already has. That rules out the Noto Serif vendored
  for the PDF; Cambria is a ClearType serif shipped with Office and named in
  the house spec.
* Headings use custom styles rather than Word's built-in Heading 1-3, whose
  blue theme colour ignores the document's own palette entirely.
* The footer is a real section footer carrying a PAGE field, suppressed on
  the title page. The previous writer appended an ordinary body paragraph,
  so its "Generated …" line appeared once, in the text flow, on whatever page
  the brief happened to end.
* Section numbers are written into the heading text instead of being driven
  by a Word list definition. That keeps them identical to the PDF's and stops
  Word renumbering them on the reader's behalf.
* Table rules are drawn per cell: none of Word's built-in table styles is
  horizontal-rules-only, and every "Light/Medium/Colorful" variant brings
  theme colours with it.
"""

import os
from datetime import date

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn
from docx.shared import Inches, Pt, RGBColor

from agent.core import ResearchResult
from services import document_theme as T
from services.document_common import (
    Block, SectionNumberer, build_references, clean, document_filename,
    normalise_headings, numeric_columns, parse_inline, parse_markdown, strip_citations,
)

H1, H2, H3 = "Veritas Heading 1", "Veritas Heading 2", "Veritas Heading 3"
HEADING_STYLES = {1: H1, 2: H2, 3: H3}
TITLE_STYLE = "Veritas Title"
SUBTITLE_STYLE = "Veritas Subtitle"
CAPTION_STYLE = "Veritas Caption"
REFERENCE_STYLE = "Veritas Reference"
CONTENTS_STYLE = "Veritas Contents"

# pPr children that must sort after w:pBdr under the OOXML schema. Word
# rejects a document whose paragraph properties are out of order, so the
# border element is inserted before these rather than appended.
_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)

# CT_Settings children that must sort after w:updateFields, for the same
# reason: python-docx's default settings.xml already ends with w:compat,
# w:rsids and w:listSeparator, so appending would land out of order.
_AFTER_UPDATE_FIELDS = (
    "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars",
    "w:rsids", "w:themeFontLang", "w:clrSchemeMapping", "w:shapeDefaults",
    "w:decimalSymbol", "w:listSeparator",
)


def _hex(color: tuple) -> str:
    return "%02X%02X%02X" % color


# --------------------------------------------------------------------------
# low-level OOXML helpers
# --------------------------------------------------------------------------

def _force_font(element, family: str):
    """
    Pins the face for every script, not just Latin.

    Setting `style.font.name` writes only w:ascii and w:hAnsi; without
    w:eastAsia and w:cs, a reader whose Word defaults differ gets a document
    in two faces.
    """
    fonts = element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), family)


def _set_outline_level(style, level: int):
    """Marks a custom style as a heading so Word's TOC field collects it."""
    marker = OxmlElement("w:outlineLvl")
    marker.set(qn("w:val"), str(level))
    style.element.get_or_add_pPr().append(marker)


def _bottom_rule(paragraph, color: tuple, size: int = 8, space: int = 4):
    """A rule beneath a paragraph — the Word twin of the PDF's accent line."""
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))        # eighths of a point
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), _hex(color))
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().insert_element_before(borders, *_AFTER_PBDR)


def _cell_rule(cell, color: tuple, size: int = 6):
    """Horizontal rule under one table cell, with no vertical lines at all."""
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "right"):
        blank = OxmlElement(f"w:{edge}")
        blank.set(qn("w:val"), "nil")
        borders.append(blank)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:color"), _hex(color))
    borders.append(bottom)
    cell._tc.get_or_add_tcPr().append(borders)


def _field(paragraph, instruction: str, placeholder: str = ""):
    """Inserts a Word field (PAGE, TOC, …) for Word to evaluate on open."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction_node, separate, text, end):
        run._r.append(node)
    return run


def _add_hyperlink(paragraph, url: str, text: str):
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)

    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), _hex(T.LINK))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)

    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)

    link.append(run)
    paragraph._p.append(link)


# --------------------------------------------------------------------------
# document setup
# --------------------------------------------------------------------------

def _define_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = T.WORD_BODY_FAMILY
    normal.font.size = Pt(T.SIZE_BODY)
    normal.font.color.rgb = RGBColor(*T.INK)
    _force_font(normal.element, T.WORD_BODY_FAMILY)
    body = normal.paragraph_format
    body.line_spacing = T.WORD_LINE_SPACING
    body.space_after = Pt(T.WORD_SPACE_AFTER)
    body.widow_control = True

    def make(name, size, *, bold=False, italic=False, color=T.INK,
             before=0, after=4, outline=None, keep=False):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = T.WORD_BODY_FAMILY
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor(*color)
        _force_font(style.element, T.WORD_BODY_FAMILY)
        layout = style.paragraph_format
        layout.space_before = Pt(before)
        layout.space_after = Pt(after)
        # keep_with_next is the spec's "never leave a heading alone at the
        # bottom of a page" — Word enforces it during its own pagination.
        layout.keep_with_next = keep
        if outline is not None:
            _set_outline_level(style, outline)
        return style

    make(H1, T.SIZE_H1, bold=True, before=16, after=5, outline=0, keep=True)
    make(H2, T.SIZE_H2, bold=True, before=12, after=4, outline=1, keep=True)
    make(H3, T.SIZE_H3, bold=True, italic=True, before=10, after=3, outline=2, keep=True)
    make(TITLE_STYLE, T.SIZE_TITLE, bold=True, after=4)
    make(SUBTITLE_STYLE, T.SIZE_SUBTITLE, italic=True, color=T.MUTED, after=2)
    make(CAPTION_STYLE, T.SIZE_CAPTION, italic=True, color=T.MUTED, before=6, after=3, keep=True)
    make(REFERENCE_STYLE, T.SIZE_REFERENCE, after=4)
    # Same look as H1 but carries no outline level, so the contents heading
    # does not list itself inside the table of contents.
    make(CONTENTS_STYLE, T.SIZE_H1, bold=True, before=0, after=5, keep=True)


def _setup_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for edge in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, edge, Inches(1))

    # Spec: page numbers only, and none on the title page.
    section.different_first_page_header_footer = True
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(footer, "PAGE", "1")
    for run in footer.runs:
        run.font.name = T.WORD_BODY_FAMILY
        run.font.size = Pt(T.SIZE_FOOTER)
        run.font.color.rgb = RGBColor(*T.MUTED)
    return section


def _set_metadata(doc, result: ResearchResult, topic: str):
    properties = doc.core_properties
    properties.title = f"{topic} — {T.DOC_KIND}"
    properties.author = T.AUTHOR
    properties.subject = T.DOC_KIND
    properties.category = T.DOC_KIND
    properties.language = "en"
    properties.keywords = ", ".join(filter(None, [topic, "research brief", result.provider]))
    properties.comments = f"Generated by {T.AUTHOR} ({result.provider})."


# --------------------------------------------------------------------------
# content
# --------------------------------------------------------------------------

def _add_runs(paragraph, text: str):
    for run in parse_inline(text):
        if not run.text:
            continue
        if run.link:
            _add_hyperlink(paragraph, run.link, run.text)
            continue
        added = paragraph.add_run(run.text)
        added.bold = run.bold          # citations stay plain, as in the PDF
        added.italic = run.italic
        if run.code:
            added.font.name = T.WORD_MONO_FAMILY
            added.font.size = Pt(T.SIZE_BODY - 1)


def _resize(paragraph, points: float):
    for run in paragraph.runs:
        run.font.size = Pt(points)


def _title_page(doc, result: ResearchResult, today: date):
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(170)
    rule.paragraph_format.space_after = Pt(4)
    _bottom_rule(rule, T.BRAND, size=18, space=1)

    doc.add_paragraph(clean(result.topic), style=TITLE_STYLE)
    doc.add_paragraph(T.DOC_KIND, style=SUBTITLE_STYLE)
    doc.add_paragraph(f"{T.AUTHOR}  ·  {today:%d %B %Y}", style=SUBTITLE_STYLE)
    doc.add_page_break()


def _add_toc(doc):
    doc.add_paragraph("Contents", style=CONTENTS_STYLE)
    _field(
        doc.add_paragraph(),
        'TOC \\o "1-3" \\h \\z \\u',
        "Contents appear here once Word updates this field (select all, then F9).",
    )
    doc.add_page_break()

    # Ask Word to refresh fields on open, so the reader normally never has to.
    # w:updateFields has a fixed position in CT_Settings; appending it lands
    # after w:compat / w:rsids / w:listSeparator, leaving settings.xml out of
    # schema order — which Word reports as a document needing repair.
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.insert_element_before(update, *_AFTER_UPDATE_FIELDS)


def _section_heading(doc, label: str, level: int):
    """One numbered heading, with the accent rule under level 1 — as in the PDF."""
    paragraph = doc.add_paragraph(label, style=HEADING_STYLES[level])
    if level == 1:
        _bottom_rule(paragraph, T.BRAND)
    return paragraph


def _add_table(doc, block: Block):
    if not block.header:
        return
    # Spec: a table's caption sits above it and reads on its own.
    if block.caption:
        doc.add_paragraph(clean(block.caption), style=CAPTION_STYLE)

    numeric = numeric_columns(block.header, block.rows)
    table = doc.add_table(rows=1, cols=len(block.header))
    table.style = doc.styles["Normal Table"]

    for cell, text, is_numeric in zip(table.rows[0].cells, block.header, numeric):
        paragraph = cell.paragraphs[0]
        paragraph.add_run(clean(text)).bold = True
        if is_numeric:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _resize(paragraph, T.SIZE_TABLE)
        _cell_rule(cell, T.INK, size=8)

    for row_data in block.rows:
        cells = table.add_row().cells
        for cell, text, is_numeric in zip(cells, row_data, numeric):
            paragraph = cell.paragraphs[0]
            _add_runs(paragraph, text)
            if is_numeric:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _resize(paragraph, T.SIZE_TABLE)
            _cell_rule(cell, T.RULE)

    # Repeat the heading row when a table breaks across pages.
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph()


def _add_references(doc, result: ResearchResult, numberer: SectionNumberer):
    _section_heading(doc, f"{numberer.next(1)}  References", 1)

    references = build_references(result.sources)
    if not references:
        doc.add_paragraph("No sources were recorded for this brief.", style=REFERENCE_STYLE)
        return

    for reference in references:
        paragraph = doc.add_paragraph(style=REFERENCE_STYLE)
        layout = paragraph.paragraph_format
        # IEEE entries hang under their own bracketed number.
        layout.left_indent = Inches(0.45)
        layout.first_line_indent = Inches(-0.45)
        paragraph.add_run(f"[{reference.index}]\t").bold = True
        paragraph.add_run(f"{reference.text} [Online]. Available: ")
        _add_hyperlink(paragraph, reference.url, reference.url)


# --------------------------------------------------------------------------
# entry points
# --------------------------------------------------------------------------

def build_docx(result: ResearchResult, today: date | None = None):
    today = today or date.today()
    topic = clean(result.topic)

    doc = Document()
    _define_styles(doc)
    _setup_page(doc)
    _set_metadata(doc, result, topic)
    _title_page(doc, result, today)

    blocks = normalise_headings(parse_markdown(strip_citations(result.markdown)))
    if len(result.markdown.split()) > T.WORD_TOC_WORD_THRESHOLD:
        _add_toc(doc)

    numberer = SectionNumberer()
    for block in blocks:
        if block.type == "heading":
            level = min(block.level, 3)
            _section_heading(doc, f"{numberer.next(level)}  {block.text}", level)
        elif block.type == "bullet":
            style = "List Number" if block.ordered else "List Bullet"
            if block.level >= 2:
                style += " 2"
            _add_runs(doc.add_paragraph(style=style), block.text)
        elif block.type == "table":
            _add_table(doc, block)
        elif block.type == "caption":
            doc.add_paragraph(clean(block.text), style=CAPTION_STYLE)
        elif block.type == "quote":
            _add_runs(doc.add_paragraph(style="Quote"), block.text)
        elif block.type == "rule":
            _bottom_rule(doc.add_paragraph(), T.RULE, size=6)
        else:
            _add_runs(doc.add_paragraph(), block.text)

    _add_references(doc, result, numberer)
    return doc


def save_research_as_docx(result: ResearchResult, output_dir: str = "output") -> str:
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, document_filename(result.topic, "docx"))
    build_docx(result).save(path)
    return path
